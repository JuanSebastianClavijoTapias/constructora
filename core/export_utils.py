"""
Módulo de exportación para reportes contables
Genera archivos Excel y PDF con información financiera consolidada
"""

from io import BytesIO
from datetime import datetime, date
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from weasyprint import HTML, CSS
from django.template.loader import render_to_string
from django.http import HttpResponse


class ExportadorExcel:
    """Genera reportes en formato Excel con estilos y formatos"""
    
    def __init__(self, titulo="Reporte Contable"):
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Reporte"
        self._setup_estilos()
        
    def _setup_estilos(self):
        """Configura estilos y formatos"""
        self.titulo_font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
        self.titulo_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        
        self.header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        
        self.currency_format = '#,##0'
        self.date_format = 'dd/mm/yyyy'
        
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def agregar_titulo(self, titulo, fila=1):
        """Agrega un título al reporte"""
        self.ws.merge_cells(f'A{fila}:H{fila}')
        cell = self.ws[f'A{fila}']
        cell.value = titulo
        cell.font = self.titulo_font
        cell.fill = self.titulo_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        self.ws.row_dimensions[fila].height = 25
        return fila + 2
    
    def agregar_tabla(self, datos, encabezados, fila=1):
        """
        Agrega una tabla con datos y encabezados
        datos: lista de diccionarios
        encabezados: lista de columnas
        """
        # Encabezados
        for col, encabezado in enumerate(encabezados, 1):
            cell = self.ws.cell(row=fila, column=col)
            cell.value = encabezado
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border
        
        # Datos
        for idx, fila_datos in enumerate(datos, fila + 1):
            for col, encabezado in enumerate(encabezados, 1):
                cell = self.ws.cell(row=idx, column=col)
                valor = fila_datos.get(encabezado, '')
                
                if isinstance(valor, Decimal) or isinstance(valor, float):
                    cell.number_format = self.currency_format
                    cell.alignment = Alignment(horizontal='right')
                elif isinstance(valor, date):
                    cell.value = valor
                    cell.number_format = self.date_format
                else:
                    cell.value = valor
                    cell.alignment = Alignment(horizontal='left')
                
                cell.border = self.border
        
        # Ajustar ancho de columnas
        for col in range(1, len(encabezados) + 1):
            self.ws.column_dimensions[chr(64 + col)].width = 15
        
        return idx + 2
    
    def obtener_archivo(self):
        """Retorna el archivo Excel como BytesIO"""
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output


class ExportadorPDF:
    """Genera reportes en formato PDF usando WeasyPrint"""
    
    @staticmethod
    def generar_reporte_html(titulo, contenido_html):
        """Crea HTML base para PDF"""
        template_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <title>{titulo}</title>
            <style>
                * {{
                    margin: 0;
                    padding: 0;
                    box-sizing: border-box;
                }}
                
                body {{
                    font-family: Arial, sans-serif;
                    font-size: 11pt;
                    color: #333;
                }}
                
                .page-break {{
                    page-break-after: always;
                }}
                
                .titulo {{
                    text-align: center;
                    font-size: 18pt;
                    font-weight: bold;
                    margin: 20px 0;
                    color: #1F4E78;
                    border-bottom: 2px solid #1F4E78;
                    padding-bottom: 10px;
                }}
                
                .fecha {{
                    text-align: right;
                    font-size: 10pt;
                    color: #666;
                    margin-bottom: 20px;
                }}
                
                table {{
                    width: 100%;
                    border-collapse: collapse;
                    margin: 15px 0;
                }}
                
                table thead {{
                    background-color: #366092;
                    color: white;
                }}
                
                table th {{
                    padding: 10px;
                    text-align: left;
                    font-weight: bold;
                    border: 1px solid #999;
                }}
                
                table td {{
                    padding: 8px 10px;
                    border: 1px solid #ddd;
                }}
                
                table tbody tr:nth-child(even) {{
                    background-color: #f5f5f5;
                }}
                
                .total-row {{
                    font-weight: bold;
                    background-color: #e8eef7;
                }}
                
                .text-right {{
                    text-align: right;
                }}
                
                .text-center {{
                    text-align: center;
                }}
                
                .resumen {{
                    margin: 20px 0;
                    padding: 15px;
                    background-color: #f5f5f5;
                    border-left: 4px solid #1F4E78;
                }}
                
                .resumen-item {{
                    display: flex;
                    justify-content: space-between;
                    margin: 8px 0;
                    padding: 5px 0;
                }}
                
                .resumen-item label {{
                    font-weight: bold;
                }}
                
                .resumen-item .valor {{
                    text-align: right;
                    min-width: 100px;
                }}
                
                .footer {{
                    margin-top: 30px;
                    padding-top: 20px;
                    border-top: 1px solid #999;
                    text-align: center;
                    font-size: 9pt;
                    color: #666;
                }}
                
                @media print {{
                    body {{
                        margin: 0;
                        padding: 10mm;
                    }}
                }}
            </style>
        </head>
        <body>
            <div class="fecha">Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}</div>
            <div class="titulo">{titulo}</div>
            {contenido_html}
            <div class="footer">
                <p>Reporte generado automáticamente por el sistema de gestión contable</p>
            </div>
        </body>
        </html>
        """
        return template_html
    
    @staticmethod
    def generar_pdf(html_content):
        """Convierte HTML a PDF usando WeasyPrint"""
        try:
            # Crear documento WeasyPrint
            from io import BytesIO
            
            output = BytesIO()
            HTML(string=html_content).write_pdf(output)
            output.seek(0)
            
            return output.getvalue()
        except Exception as e:
            raise Exception(f"Error generando PDF: {str(e)}")


def exportar_nominas_excel(nominas, propiedad):
    """Exporta nóminas a Excel"""
    exportador = ExportadorExcel(f"Nóminas - {propiedad.nombre}")
    
    fila = exportador.agregar_titulo(f"NÓMINAS DEL MES - {propiedad.nombre}")
    
    datos = []
    total_neto = Decimal('0')
    
    for nomina in nominas:
        total_neto += nomina.total_neto
        datos.append({
            'Empleado': f"{nomina.empleado.nombres} {nomina.empleado.apellidos}",
            'Cédula': nomina.empleado.cedula,
            'Cargo': nomina.empleado.cargo,
            'Salario Base': float(nomina.salario_base),
            'Deducciones': float(nomina.total_descuentos),
            'Aportes Patronales': float(nomina.total_aportes_patronales),
            'Neto a Pagar': float(nomina.total_neto),
            'Estado': nomina.estado
        })
    
    encabezados = ['Empleado', 'Cédula', 'Cargo', 'Salario Base', 'Deducciones',
                   'Aportes Patronales', 'Neto a Pagar', 'Estado']
    exportador.agregar_tabla(datos, encabezados, fila)
    
    # Agregar total
    ultima_fila = fila + len(datos) + 1
    cell = exportador.ws[f'A{ultima_fila}']
    cell.value = "TOTAL NETO"
    cell.font = Font(bold=True)

    cell = exportador.ws[f'G{ultima_fila}']
    cell.value = float(total_neto)
    cell.font = Font(bold=True)
    cell.number_format = exportador.currency_format
    cell.fill = PatternFill(start_color='FFFF99', end_color='FFFF99', fill_type='solid')

    return exportador.obtener_archivo()


def exportar_obligaciones_excel(obligaciones, propiedad):
    """Exporta obligaciones fiscales a Excel"""
    exportador = ExportadorExcel(f"Obligaciones Fiscales - {propiedad.nombre}")
    
    fila = exportador.agregar_titulo(f"OBLIGACIONES FISCALES - {propiedad.nombre}")
    
    datos = []
    total = Decimal('0')
    
    for obligacion in obligaciones:
        total += obligacion.monto_obligacion
        datos.append({
            'Obligación': obligacion.get_tipo_obligacion_display(),
            'Descripción': obligacion.descripcion or '-',
            'Monto': float(obligacion.monto_obligacion),
            'Vencimiento': obligacion.fecha_vencimiento_proximo,
            'Frecuencia': obligacion.get_frecuencia_pago_display(),
            'Activa': 'Sí' if obligacion.activa else 'No',
            'Referencia': obligacion.referencia_externa or '-'
        })
    
    encabezados = ['Obligación', 'Descripción', 'Monto', 'Vencimiento', 'Frecuencia', 'Activa', 'Referencia']
    exportador.agregar_tabla(datos, encabezados, fila)
    
    # Agregar total
    ultima_fila = fila + len(datos) + 1
    cell = exportador.ws[f'A{ultima_fila}']
    cell.value = "TOTAL OBLIGACIONES"
    cell.font = Font(bold=True)
    
    cell = exportador.ws[f'C{ultima_fila}']
    cell.value = float(total)
    cell.font = Font(bold=True)
    cell.number_format = exportador.currency_format
    
    return exportador.obtener_archivo()


def exportar_reporte_mensual_excel(reporte, nominas, obligaciones, servicios, propiedad):
    """Exporta reporte mensual consolidado a Excel"""
    exportador = ExportadorExcel()
    
    # Título
    fila = exportador.agregar_titulo(f"REPORTE CONTABLE MENSUAL - {propiedad.nombre}")
    fila += 1
    
    # Resumen general
    resumen_datos = [{
        'Concepto': 'TOTAL NÓMINA',
        'Porcentaje': '-',
        'Monto': float(reporte.total_nomina),
        'Observaciones': 'Salarios brutos'
    }, {
        'Concepto': 'APORTES PATRONALES',
        'Porcentaje': '-',
        'Monto': float(reporte.total_aportes_patronales),
        'Observaciones': 'Pensión, Salud, SENA, ICBF, Caja, ARL'
    }, {
        'Concepto': 'OBLIGACIONES FISCALES',
        'Porcentaje': '-',
        'Monto': float(reporte.total_obligaciones),
        'Observaciones': 'Impuestos y aportes'
    }, {
        'Concepto': 'SERVICIOS PÚBLICOS',
        'Porcentaje': '-',
        'Monto': float(reporte.total_servicios_publicos),
        'Observaciones': 'Agua, luz, gas, internet'
    }, {
        'Concepto': 'MANTENIMIENTO',
        'Porcentaje': '-',
        'Monto': float(reporte.total_mantenimiento),
        'Observaciones': 'Contratos y reparaciones'
    }, {
        'Concepto': 'TOTAL GASTOS',
        'Porcentaje': '-',
        'Monto': float(reporte.total_gastos),
        'Observaciones': 'Todos los gastos consolidados'
    }]
    
    encabezados = ['Concepto', 'Porcentaje', 'Monto', 'Observaciones']
    fila = exportador.agregar_tabla(resumen_datos, encabezados, fila)
    
    return exportador.obtener_archivo()


def exportar_reporte_pdf(titulo, contexto):
    """Crea un response HTTP con PDF"""
    # Renderizar template HTML
    html_content = render_to_string('core/contabilidad/reporte_pdf.html', contexto)
    
    # Generar PDF
    html_completo = ExportadorPDF.generar_reporte_html(titulo, html_content)
    pdf_bytes = ExportadorPDF.generar_pdf(html_completo)
    
    # Crear response
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{titulo}.pdf"'
    
    return response
